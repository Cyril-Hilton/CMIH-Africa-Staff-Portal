"""Read-only source extraction. Retain source values/formulas without evaluating spreadsheet code."""
import argparse,datetime,json,hashlib
import openpyxl
p=argparse.ArgumentParser();p.add_argument('workbook');p.add_argument('output');args=p.parse_args()
v=openpyxl.load_workbook(args.workbook,data_only=True);f=openpyxl.load_workbook(args.workbook,data_only=False)
rows=[]
for s in v:
 for row in s:
  cells={};formulas={}
  for c in row:
   raw=f[s.title][c.coordinate]
   if c.value is not None:cells[c.column_letter]=c.value.isoformat() if isinstance(c.value,(datetime.datetime,datetime.date)) else c.value
   if raw.data_type=='f':formulas[c.column_letter]=raw.value
  if not cells and not formulas:continue
  payload={'sheet':s.title,'source_row':row[0].row,'hidden_sheet':s.sheet_state!='visible','cells':cells,'formulas':formulas}
  payload['checksum']=hashlib.sha256(json.dumps(payload,sort_keys=True,ensure_ascii=False).encode()).hexdigest();rows.append(payload)
with open(args.output,'w') as out:json.dump({'source_id':'1ydsMLK86iDG5XPp1PcC3_qjh_2CjeJ_-','rows':rows},out,ensure_ascii=False)
print(f'Extracted {len(rows)} nonempty source rows from {len(v.sheetnames)} tabs; source unchanged.')
