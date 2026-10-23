"""Copy labelled account allocation fields; do not invent a month, period or realized revenue."""
import openpyxl,json,argparse,re
p=argparse.ArgumentParser();p.add_argument('workbook');p.add_argument('output');a=p.parse_args();w=openpyxl.load_workbook(a.workbook,data_only=True);rows=[]
def norm(v):return re.sub(r'\s+',' ',str(v or '').strip()).upper()
def num(v):
 if isinstance(v,bool) or v is None:return None
 if isinstance(v,(int,float)):return float(v)
 s=re.sub(r'[,\s]','',str(v));return float(s) if re.fullmatch(r'-?\d+(\.\d+)?',s) else None
active=['2026 ALFAH BARAKAH Pipeline','2026 Nicolette Nana Ekua Pipeli','2026 MASANA ALFAH  Pipeline','2026 JOVITA ANANKUM pipeline']
for name in active[:2]:
 s=w[name];header=None
 for row in s:
  values={c.column:norm(c.value) for c in row if c.value is not None}
  if 'BRANDS/ACCOUNT' in values.values() and 'YTD TARGET' in values.values():header={v:k for k,v in values.items() if v in ['BRANDS/ACCOUNT','YTD TARGET','MONTHLY TARGET']};continue
  if not header:continue
  client=s.cell(row[0].row,header['BRANDS/ACCOUNT']).value
  if not isinstance(client,str) or not client.strip():continue
  ytd=num(s.cell(row[0].row,header['YTD TARGET']).value);monthly=num(s.cell(row[0].row,header['MONTHLY TARGET']).value) if 'MONTHLY TARGET' in header else None
  if ytd is None and monthly is None:continue
  rows.append({'sheet':name,'row':row[0].row,'client':client.strip(),'owner_sheet':name,'ytd_target':ytd,'monthly_target':monthly,'budget_target':None,'reported_target':None,'source_values':{s.cell(row[0].row,c).column_letter:s.cell(row[0].row,c).value for c in header.values()}})
# This allocation block has a budget column but no annual/monthly heading: retain its own label.
s=w[active[2]]
for row in s:
 if row[0].row>25:break
 if isinstance(row[0].value,(int,float)) and isinstance(row[1].value,str) and norm(row[5].value)=='MASANA ALFAH':
  value=num(row[6].value)
  if value is not None:rows.append({'sheet':s.title,'row':row[0].row,'client':row[1].value.strip(),'owner_sheet':s.title,'ytd_target':None,'monthly_target':None,'budget_target':value,'reported_target':None,'source_values':{'B':row[1].value,'G':row[6].value}})
# Separate master allocation tab. Client Name is the contact person; Activities / Brands is the account.
s=w['Total Client Allocation '];aliases={'NICOLETTE':active[1],'MASANA':active[2],'MASANA ALFAH':active[2],'ALFAH BARAKAH':active[0],'JOVITA':active[3],'JOVITA ANANKUM':active[3]};unmapped=[]
for row in s:
 if not isinstance(row[0].value,(int,float)) or not isinstance(row[1].value,str):continue
 owner=aliases.get(norm(row[5].value));budget=num(row[6].value);target=num(row[7].value)
 if budget is None and target is None:continue
 if not owner:unmapped.append(row[0].row)
 rows.append({'sheet':s.title,'row':row[0].row,'client':row[1].value.strip(),'owner_sheet':owner,'reported_owner':str(row[5].value or '').strip(),'ytd_target':None,'monthly_target':None,'budget_target':budget,'reported_target':target,'source_values':{'B':row[1].value,'G':row[6].value,'H':row[7].value}})
json.dump({'source_id':'1ydsMLK86iDG5XPp1PcC3_qjh_2CjeJ_-','rows':rows,'unmapped_owner_rows':unmapped},open(a.output,'w'),ensure_ascii=False);print(json.dumps({'allocations':len(rows),'unmapped_owner_rows':unmapped}))
