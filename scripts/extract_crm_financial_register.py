"""Read the maintained workbook's separate PO/invoice registers. No revenue or cash is inferred."""
import argparse,hashlib,json,re,datetime
import openpyxl
p=argparse.ArgumentParser();p.add_argument('workbook');p.add_argument('output');a=p.parse_args()
w=openpyxl.load_workbook(a.workbook,data_only=True);out=[]
active=['2026 ALFAH BARAKAH Pipeline','2026 Nicolette Nana Ekua Pipeli','2026 MASANA ALFAH  Pipeline','2026 JOVITA ANANKUM pipeline']
def num(v):
 if isinstance(v,bool) or v is None:return None
 if isinstance(v,(int,float)):return float(v)
 s=str(v).strip().replace(',','')
 return float(s) if re.fullmatch(r'-?\d+(\.\d+)?',s) else None
for name in active:
 s=w[name];header=None;lastclient=''
 for row in s:
  values={c.column:str(c.value).strip().upper() for c in row if c.value is not None}
  if 'PO RAISED' in values.values() and 'INVOICED' in values.values():
   header={label:col for col,label in values.items() if label in ['BRAND','PROJECT','PO RAISED','PO BALANCE','STATUS','INVOICED','INVOICE STATUS']};lastclient='';continue
  if not header:continue
  if any(v in ['BRANDS/ACCOUNT','BRANDS/ACCOUNTS','YTD TARGET','MONTHLY TARGET'] for col,v in values.items() if min(header.values())<=col<=max(header.values())):break
  def val(key):return s.cell(row[0].row,header[key]).value if key in header else None
  client=str(val('BRAND') or '').strip();project=str(val('PROJECT') or '').strip()
  if not project or project.upper() in ['PROJECT','TOTAL','MONTHLY TARGET']:continue
  if client and re.fullmatch(r'\d+(\.\d+)?',client):continue
  if client:lastclient=client
  client=client or lastclient
  if not client:continue
  po,balance,invoice=num(val('PO RAISED')),num(val('PO BALANCE')),num(val('INVOICED'))
  if all(x is None for x in [po,balance,invoice]):continue
  source={s.cell(row[0].row,col).column_letter:s.cell(row[0].row,col).value for col in header.values()}
  source={k:v.isoformat() if isinstance(v,(datetime.date,datetime.datetime)) else v for k,v in source.items()}
  checksum=hashlib.sha256(json.dumps(source,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
  out.append({'sheet':name,'row':row[0].row,'client':client,'project':project,'po_amount':po,'po_balance':balance,'invoiced_amount':invoice,'po_status':str(val('STATUS') or '').strip() or None,'invoice_status':str(val('INVOICE STATUS') or '').strip() or None,'source_year':2026,'as_of':None,'source_values':source,'checksum':checksum})
json.dump({'source_id':'1ydsMLK86iDG5XPp1PcC3_qjh_2CjeJ_-','rows':out},open(a.output,'w'),ensure_ascii=False)
print(json.dumps({'financial_register_rows':len(out),'po_total':sum(r['po_amount'] or 0 for r in out),'invoiced_total':sum(r['invoiced_amount'] or 0 for r in out),'undated_register_rows':len(out)}))
